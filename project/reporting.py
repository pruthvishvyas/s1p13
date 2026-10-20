import os
import base64
import datetime
import numpy as np
import pandas as pd
from utils import *  # @module-import
from features import channel_table  # @module-import


def final_summary(m, ctx, cfg):
    c = cfg
    act = m[m[c['STATUS']] == c['ACTIVE_VALUE']]
    lab = m[m['is_churned'].notna()]
    ret = ctx['goals']['retention']
    mc = ctx.get('model_card') or {}
    mt = mc.get('metrics_holdout', {})
    ps = ctx['logic']['priority_summary'].set_index('priority')
    s = {'ARR': float(act[c['ARR']].sum()), 'MRR': float(act[c['MRR']].sum()), 'active_accounts': int(len(act)),
         'total_accounts': int(len(m)), 'logo_churn_pct': float((lab['is_churned'] == 1).mean() * 100),
         'revenue_churn_pct': float(lab.loc[lab['is_churned'] == 1, c['MRR0']].sum() / max(lab[c['MRR0']].sum(), 1e-9) * 100),
         'GRR_pct': ret['grr_pct'], 'NRR_pct': ret['nrr_pct'], 'ARR_at_risk': float(act['arr_at_risk'].sum()),
         'P1_count': int((act['priority'] == 'P1').sum()), 'P1_P2_arr_at_risk': float(ctx['logic']['headline_arr_at_risk']),
         'model_pr_auc': mt.get('pr_auc'), 'model_roc_auc': mt.get('roc_auc'), 'scoring_source': mc.get('primary_scoring', 'rule'),
         'forecast_label': (ctx.get('fc') or {}).get('meta', {}).get('label', 'n/a')}
    print('=' * 60)
    print('FINAL SUMMARY')
    for k, v in s.items():
        print('  %-20s %s' % (k, round(v, 2) if isinstance(v, float) else v))
    print('=' * 60)
    return s


def build_excel(m, ctx, cfg, summary):
    c = cfg
    path = c['EXCEL']
    os.makedirs(os.path.dirname(path), exist_ok=True)
    act = m[m[c['STATUS']] == c['ACTIVE_VALUE']]
    lab = m[m['is_churned'].notna()]
    plan = lab.groupby(c['PLAN']).agg(accounts=(c['ID'], 'size'), churn_rate_pct=('is_churned', lambda x: x.mean() * 100),
                                      avg_initial_mrr=(c['MRR0'], 'mean'), avg_discount_pct=(c['DISCOUNT'], 'mean'),
                                      median_payback_months=(c['PAYBACK'], 'median')).reset_index()
    plan.loc[plan['accounts'] < c['MIN_GROUP_N'], 'churn_rate_pct'] = np.nan
    summ = pd.DataFrame({'Metric': list(summary.keys()), 'Value': [round(v, 2) if isinstance(v, float) else v for v in summary.values()]})
    save = act[act['priority'].isin(['P1', 'P2'])].sort_values(['priority', 'arr_at_risk'], ascending=[True, False])
    save = save[[c['ID'], c['COMPANY'], c['PLAN'], 'priority', 'churn_prob', 'churn_band', c['ARR'], 'arr_at_risk', 'recommendation', 'drivers']].head(100)
    coll = act[(act['late_payer_flag'] == 1) | (act[c['BALANCE']] > 0)].sort_values(c['BALANCE'], ascending=False)
    coll = coll[[c['ID'], c['COMPANY'], c[ 'DAYS_PAY'], c['LATE'], c['BALANCE'], 'balance_to_mrr', c['ARR'], 'churn_prob']].head(100)
    seg = ctx['seg']['summary'] if len(ctx['seg'].get('summary', [])) else pd.DataFrame({'note': ['segmentation skipped']})
    fc = ctx['fc']['monthly'] if ctx.get('fc') and len(ctx['fc'].get('monthly', [])) else pd.DataFrame({'note': ['forecast skipped']})
    sheets = {'Revenue Summary': summ, 'Plan & Channel Economics': plan, 'Account Segments': seg,
              'Save List (This Week)': save, 'Collections': coll, 'Forecast': fc}
    with pd.ExcelWriter(path, engine='openpyxl') as xw:
        summ.to_excel(xw, sheet_name='Revenue Summary', index=False)
        plan.to_excel(xw, sheet_name='Plan & Channel Economics', index=False)
        ch = ctx['logic']['channel_recs']
        ch.to_excel(xw, sheet_name='Plan & Channel Economics', index=False, startrow=len(plan) + 3)
        for nm in ['Account Segments', 'Save List (This Week)', 'Collections', 'Forecast']:
            sheets[nm].to_excel(xw, sheet_name=nm, index=False)
        for ws in xw.book.worksheets:
            for col in ws.columns:
                width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col[:60])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
            ws.freeze_panes = 'A2'
    print('Excel report saved:', path)
    return path


def build_dashboard(m, ctx, cfg):
    c = cfg
    try:
        import plotly.graph_objects as go
        import plotly.io as pio
    except ImportError:
        print('plotly not installed -> interactive dashboard skipped (pip install plotly).')
        return None
    act = m[m[c['STATUS']] == c['ACTIVE_VALUE']]
    figs = []
    f1 = go.Figure(go.Histogram(x=act['churn_prob'], nbinsx=20))
    f1.update_layout(title='Active accounts by calibrated churn probability', xaxis_title='churn probability')
    figs.append(f1)
    ps = ctx['logic']['priority_summary']
    figs.append(go.Figure(go.Bar(x=ps['priority'], y=ps['arr_at_risk'])).update_layout(title='ARR at risk by priority'))
    mf = ctx['fc']['monthly'] if ctx.get('fc') and len(ctx['fc'].get('monthly', [])) else None
    if mf is not None:
        f3 = go.Figure()
        f3.add_trace(go.Scatter(x=mf['month'], y=mf['mrr_best'], mode='lines', name='best', line=dict(width=0)))
        f3.add_trace(go.Scatter(x=mf['month'], y=mf['mrr_worst'], mode='lines', name='worst', fill='tonexty', line=dict(width=0)))
        f3.add_trace(go.Scatter(x=mf['month'], y=mf['mrr_base'], mode='lines+markers', name='base'))
        f3.update_layout(title='MRR run-off forecast')
        figs.append(f3)
    ch = ctx['logic']['channel_recs']
    figs.append(go.Figure(go.Bar(x=ch[c['CHANNEL']].astype(str), y=ch['ltv_cac'])).update_layout(title='Realized LTV:CAC by channel'))
    top = act.sort_values('arr_at_risk', ascending=False).head(15)
    figs.append(go.Figure(go.Bar(x=top['arr_at_risk'], y=top[c['COMPANY']].astype(str), orientation='h')).update_layout(title='Top 15 accounts by ARR at risk', yaxis=dict(autorange='reversed')))
    html = ['<html><head><meta charset="utf-8"><title>SaaS Revenue Retention Dashboard</title></head><body style="font-family:Arial;max-width:1000px;margin:auto"><h1>SaaS Revenue Retention Dashboard</h1>']
    for i, f in enumerate(figs):
        html.append(pio.to_html(f, full_html=False, include_plotlyjs=True if i == 0 else False))
    html.append('</body></html>')
    path = os.path.join(c['REPORTS_DIR'], 'dashboard.html')
    with open(path, 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(html))
    print('Dashboard saved:', path)
    return path


def _slides(m, ctx, cfg, s):
    c = cfg
    mc = ctx.get('model_card') or {}
    mt = mc.get('metrics_holdout', {})
    vr = load_json('data/processed/validation_report.json', {}) or {}
    ins = load_json(os.path.join(c['REPORTS_DIR'], 'insights.json'), []) or []
    act = m[m[c['STATUS']] == c['ACTIVE_VALUE']]
    top = act.sort_values('arr_at_risk', ascending=False).head(5)
    r = c['REPORTS_DIR']
    fmt = lambda v: '{:,.0f}'.format(v)
    sl = [
        ('Which accounts will churn, and what do we do this week?', ['SaaS revenue-retention analysis', 'Generated %s' % datetime.date.today().isoformat()], None),
        ('Executive summary', ['Active ARR %s (MRR %s) across %d active accounts' % (fmt(s['ARR']), fmt(s['MRR']), s['active_accounts']),
                               'Expected ARR at risk: %s; P1+P2 exposure %s' % (fmt(s['ARR_at_risk']), fmt(s['P1_P2_arr_at_risk'])),
                               '%d P1 accounts need action this week' % s['P1_count'],
                               'GRR %.1f%% / NRR %.1f%% (snapshot approximation)' % (s['GRR_pct'], s['NRR_pct'])], None),
        ('Business question and scope', ['Which active accounts are most likely to churn in 90 days?', 'How much ARR is at risk and what should we do for each?',
                                         'Data: one row per account, snapshot, %d rows' % vr.get('rows', len(m))], None),
        ('Data health', ['Rows %s, columns %s' % (vr.get('rows'), vr.get('cols')), 'Integrity checks: ' + ', '.join('%s=%s' % (k, v) for k, v in list((vr.get('integrity_ok') or {}).items())[:6]),
                         'Churn rate %.1f%%' % (vr.get('target_balance', {}).get('churn_rate', 0) * 100)], None),
        ('Revenue: new, expansion, contraction, churn', ['Waterfall of MRR by component', 'Logo churn %.1f%%, revenue churn %.1f%%' % (s['logo_churn_pct'], s['revenue_churn_pct'])], r + '01_mrr_waterfall_cohorts.png'),
        ('Where churn concentrates', ['Churn rate by plan, industry and region (groups with n<%d suppressed)' % c['MIN_GROUP_N']], r + '02_churn_by_segment.png'),
        ('Why accounts leave', ['Churn reasons weighted by MRR lost'], r + '03_churn_reasons_pareto.png'),
        ('Channel unit economics', ['LTV:CAC and payback by acquisition channel vs 3.0x / 12-month benchmarks'], r + '04_channel_economics.png'),
        ('Discounting', ['Discount depth versus churn and margin per seat'], r + '05_discount_churn_margin.png'),
        ('Churn model quality', ['Hold-out PR-AUC %.2f vs base rate %.2f; ROC-AUC %.2f' % (mt.get('pr_auc', float('nan')), mt.get('base_rate_test', float('nan')), mt.get('roc_auc', float('nan'))),
                                 'Scoring source: %s; leakage columns excluded' % s['scoring_source']], r + '05_churn_model_diagnostics.png'),
        ('Save list: top accounts by ARR at risk', ['%s: %s at risk (%.0f%%)' % (x[c['COMPANY']], fmt(x['arr_at_risk']), x['churn_prob'] * 100) for _, x in top.iterrows()], None),
        ('Account segments', ['Segments with plain-English traits; anomalies on a watch list'], r + '06_segments_pca.png'),
        ('6-month MRR run-off forecast', ['Forecast label: %s' % s['forecast_label'], 'Bands: 10th-90th percentile of bootstrap'], r + '07_mrr_forecast.png'),
        ('Recommendations and next steps', [i['action'] for i in ins[:4]] or ['Review the save list'], None)]
    return sl


def build_report_deck(m, ctx, cfg, summary):
    c = cfg
    slides = _slides(m, ctx, cfg, summary)
    html = ['<html><head><meta charset="utf-8"><title>Report</title><style>body{font-family:Arial;background:#eee}.s{background:#fff;width:960px;min-height:540px;margin:20px auto;padding:30px;box-shadow:0 2px 6px #999;page-break-after:always}h1{font-size:28px;color:#1a365d}li{font-size:18px;margin:6px 0}img{max-width:100%;max-height:360px}</style></head><body>']
    for i, (t, bl, img) in enumerate(slides, 1):
        html.append('<div class="s"><h1>%d. %s</h1><ul>%s</ul>' % (i, t, ''.join('<li>%s</li>' % b for b in bl)))
        if img and os.path.exists(img):
            html.append('<img src="data:image/png;base64,%s"/>' % base64.b64encode(open(img, 'rb').read()).decode('ascii'))
        html.append('</div>')
    html.append('</body></html>')
    hp = os.path.join(c['REPORTS_DIR'], c['CLIENT_ID'] + '_report.html')
    with open(hp, 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(html))
    out = [hp]
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
        for t, bl, img in slides:
            sld = prs.slides.add_slide(prs.slide_layouts[6])
            tb = sld.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(1)).text_frame
            tb.text = t
            tb.paragraphs[0].font.size = Pt(30)
            tb.paragraphs[0].font.bold = True
            bx = sld.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(5.6 if img and os.path.exists(img) else 12), Inches(5)).text_frame
            bx.word_wrap = True
            for j, b in enumerate(bl):
                p = bx.paragraphs[0] if j == 0 else bx.add_paragraph()
                p.text = '- ' + str(b)
                p.font.size = Pt(16)
            if img and os.path.exists(img):
                sld.shapes.add_picture(img, Inches(6.4), Inches(1.4), width=Inches(6.5))
        pp = os.path.join(c['REPORTS_DIR'], c['CLIENT_ID'] + '_report.pptx')
        prs.save(pp)
        out.append(pp)
    except ImportError:
        print('python-pptx not installed -> HTML deck only.')
    print('Report deck (%d slides):' % len(slides), out)
    return out


def write_readme(summary, cfg):
    c = cfg
    txt = """# SaaS Revenue-Retention Pipeline ({cid})

Question: which active accounts will churn in the next 90 days, how much ARR is at risk, and what do we do this week?

## Run
1. Put the data at `{raw}` (UTF-8 CSV, one row per account).
2. `python generate_notebook.py` then run `{cid}_pipeline.ipynb` top to bottom (or `python project/main.py`).
3. Outputs: `reports/` (charts, Excel, dashboard.html, deck), `models/`, `data/processed/`, `output/` (frontend JSON), `config/goals.json`.

## Headline (last run)
- Active ARR {arr:,.0f}; ARR at risk {risk:,.0f}; P1 accounts {p1}; scoring source: {src}.

## Method notes
- Leakage columns (Status, Churn_Date, Churn_Reason, Months_Active, lifetime metrics, tenure) are never used by the classifier.
- Health_Score / Churn_Risk are tested for leakage and excluded when they look derived from the outcome.
- Active accounts are scored with out-of-fold calibrated probabilities (never by a model that saw them).
- Probabilities rank relative risk; the 90-day horizon is an assumption because the data is a snapshot.
- GRR/NRR are snapshot approximations, not annualized.
- Thresholds live in CONFIG['THRESHOLDS'] and `data/processed/validation_report.json`; goals in `config/goals.json`.
- BigQuery: fill CONFIG['BIGQUERY'] and upload `data/processed/account_master.csv`; `export_for_frontend.py` falls back to CSV on any error.
""".format(cid=c['CLIENT_ID'], raw=c['RAW_CSV'], arr=summary['ARR'], risk=summary['ARR_at_risk'], p1=summary['P1_count'], src=summary['scoring_source'])
    with open('README.md', 'w', encoding='utf-8') as fh:
        fh.write(txt)
    print('README.md written')
